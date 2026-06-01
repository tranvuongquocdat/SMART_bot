"""DocumentExtractor — PDF / DOCX / XLSX / TXT.

Dispatch is by ``content_type`` (mime) or, if absent, by URL extension.
Each branch is independently importable so a missing dependency degrades
to an empty result rather than crashing the whole adapter import.
"""

from __future__ import annotations

import io
import logging
from typing import Any

from src.media.base import MediaExtractResult
from src.media.registry import media_adapter

log = logging.getLogger(__name__)

MAX_PAGES = 20


@media_adapter(supports={"pdf", "docx", "xlsx", "txt"})
class DocumentExtractor:
    async def extract(
        self,
        url: str | None = None,
        content: bytes | None = None,
        content_type: str | None = None,
    ) -> MediaExtractResult:
        kind = _detect_kind(content_type, url)
        if not kind or content is None:
            return MediaExtractResult(media_text="")
        try:
            if kind == "pdf":
                return _extract_pdf(content)
            if kind == "docx":
                return _extract_docx(content)
            if kind == "xlsx":
                return _extract_xlsx(content)
            if kind == "txt":
                return _extract_txt(content)
        except Exception:
            log.exception("document extract failed kind=%s", kind)
        return MediaExtractResult(media_text="")


def _detect_kind(content_type: str | None, url: str | None) -> str | None:
    ct = (content_type or "").lower()
    if "pdf" in ct:
        return "pdf"
    if "wordprocessingml" in ct or "msword" in ct:
        return "docx"
    if "spreadsheetml" in ct or "ms-excel" in ct:
        return "xlsx"
    if ct.startswith("text/"):
        return "txt"
    if url:
        u = url.lower()
        for ext in ("pdf", "docx", "xlsx", "txt"):
            if u.endswith("." + ext):
                return ext
    return None


def _extract_pdf(content: bytes) -> MediaExtractResult:
    import pypdf  # type: ignore

    reader = pypdf.PdfReader(io.BytesIO(content))
    parts: list[str] = []
    for i, page in enumerate(reader.pages):
        if i >= MAX_PAGES:
            break
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            log.warning("pdf page %d extract failed", i)
    title = None
    try:
        meta: Any = reader.metadata
        if meta is not None:
            title = getattr(meta, "title", None) or meta.get("/Title")  # type: ignore[union-attr]
    except Exception:
        pass
    return MediaExtractResult(
        media_text="\n\n".join(p for p in parts if p),
        title=str(title) if title else None,
        extra={"pages": min(len(reader.pages), MAX_PAGES)},
    )


def _extract_docx(content: bytes) -> MediaExtractResult:
    import docx  # type: ignore

    doc = docx.Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs if p.text]
    return MediaExtractResult(media_text="\n".join(parts))


def _extract_xlsx(content: bytes) -> MediaExtractResult:
    import openpyxl  # type: ignore

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets[:5]:
        lines.append(f"## {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                lines.append("\t".join(cells))
    return MediaExtractResult(media_text="\n".join(lines))


def _extract_txt(content: bytes) -> MediaExtractResult:
    text = content.decode("utf-8", errors="ignore")
    return MediaExtractResult(media_text=text)
